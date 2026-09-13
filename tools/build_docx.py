#!/usr/bin/env python3
"""build_docx.py - CONTESTED as a Word research note."""
import os, sys
from docx import Document
from docx.shared import Pt, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIGS = os.path.join(ROOT, "figs")
OUT  = os.path.join(ROOT, "CONTESTED.docx")

INK="101010"; INK2="3D3C38"; INK3="6B6A63"; ACC="C04A12"; LINE="DEDCD6"; PANEL="F6F5F2"

doc = Document()
for s in doc.sections:
    s.page_width, s.page_height = Inches(8.27), Inches(11.69)
    s.left_margin = s.right_margin = Inches(0.72)
    s.top_margin = Inches(0.66); s.bottom_margin = Inches(0.62)

st = doc.styles["Normal"]
st.font.name = "Calibri"; st.font.size = Pt(9.5); st.font.color.rgb = RGBColor.from_string(INK)
st.paragraph_format.space_after = Pt(5); st.paragraph_format.line_spacing = 1.12
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")

def shade(cell, hex_):
    el = OxmlElement("w:shd"); el.set(qn("w:val"),"clear"); el.set(qn("w:fill"),hex_)
    cell._tc.get_or_add_tcPr().append(el)

def borders(cell, **kw):
    tcPr = cell._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for edge,(sz,col) in kw.items():
        e = OxmlElement(f"w:{edge}"); e.set(qn("w:val"),"single")
        e.set(qn("w:sz"),str(sz)); e.set(qn("w:color"),col); b.append(e)
    tcPr.append(b)

def rule(width=12, color=INK, space=4):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(space); p.paragraph_format.space_after=Pt(space)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); bot = OxmlElement("w:bottom")
    bot.set(qn("w:val"),"single"); bot.set(qn("w:sz"),str(width)); bot.set(qn("w:color"),color); bot.set(qn("w:space"),"1")
    pbdr.append(bot); pPr.append(pbdr); return p

def para(text="", size=9.5, bold=False, italic=False, color=INK, align=None,
         before=0, after=5, indent=0):
    p = doc.add_paragraph(); pf = p.paragraph_format
    pf.space_before=Pt(before); pf.space_after=Pt(after)
    if indent: pf.left_indent = Inches(indent)
    if align: p.alignment = align
    else: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, text, size, bold, italic, color); return p

def add_runs(p, text, size=9.5, bold=False, italic=False, color=INK):
    """**bold** and *italic* inline."""
    import re
    for tok in re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*)", text):
        if not tok: continue
        b, i, t = bold, italic, tok
        if tok.startswith("**") and tok.endswith("**"): b, t = True, tok[2:-2]
        elif tok.startswith("*") and tok.endswith("*"): i, t = True, tok[1:-1]
        r = p.add_run(t); r.font.size=Pt(size); r.bold=b; r.italic=i
        r.font.color.rgb = RGBColor.from_string(color); r.font.name="Calibri"

def h1(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(15); p.paragraph_format.space_after=Pt(3)
    r = p.add_run(text); r.font.size=Pt(14); r.bold=True; r.font.color.rgb=RGBColor.from_string(INK)
    r.font.name="Calibri"
    pPr=p._p.get_or_add_pPr(); pb=OxmlElement("w:pBdr"); bt=OxmlElement("w:bottom")
    bt.set(qn("w:val"),"single"); bt.set(qn("w:sz"),"10"); bt.set(qn("w:color"),INK); bt.set(qn("w:space"),"3")
    pb.append(bt); pPr.append(pb)
    kn=OxmlElement("w:keepNext"); pPr.append(kn)

def h2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(11); p.paragraph_format.space_after=Pt(2)
    r = p.add_run(text); r.font.size=Pt(10.5); r.bold=True; r.font.color.rgb=RGBColor.from_string(ACC)
    r.font.name="Calibri"
    OxmlElement("w:keepNext"); p._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))

def label(kind, num, title):
    p = doc.add_paragraph(); p.paragraph_format.space_before=Pt(11); p.paragraph_format.space_after=Pt(2)
    r = p.add_run(f"{kind} {num}"); r.font.size=Pt(8); r.bold=True
    r.font.color.rgb=RGBColor.from_string(ACC); r.font.name="Calibri"
    r2 = p.add_run("   "+title); r2.font.size=Pt(9.5); r2.bold=True
    r2.font.color.rgb=RGBColor.from_string(INK); r2.font.name="Calibri"
    p._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))

def source(text):
    para(text, size=7.6, italic=True, color=INK3, align=WD_ALIGN_PARAGRAPH.LEFT, after=8)

def table(rows, widths=None, align=None, header=True, size=8.2, first_bold=False):
    t = doc.add_table(rows=0, cols=len(rows[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            p.paragraph_format.space_before=Pt(1.6); p.paragraph_format.space_after=Pt(1.6)
            a = (align[ci] if align and ci < len(align) else "l")
            p.alignment = {"l":WD_ALIGN_PARAGRAPH.LEFT,"r":WD_ALIGN_PARAGRAPH.RIGHT,
                           "c":WD_ALIGN_PARAGRAPH.CENTER}[a]
            hdr = header and ri==0
            add_runs(p, str(val), size=size-0.3 if hdr else size,
                     bold=hdr or (first_bold and ci==0), color=INK2 if hdr else INK)
            if hdr:
                shade(cells[ci], PANEL)
                borders(cells[ci], top=(10,INK), bottom=(5,INK3))
            else:
                borders(cells[ci], bottom=(3,LINE) if ri<len(rows)-1 else (10,INK))
            if widths: cells[ci].width = Inches(widths[ci])
    return t

def fig(name, width=6.6):
    p = doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before=Pt(3); p.paragraph_format.space_after=Pt(2)
    p.add_run().add_picture(os.path.join(FIGS,name), width=Inches(width))

def pagebreak(): doc.add_page_break()

def callout(text, color=ACC):
    t = doc.add_table(rows=1, cols=1); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    c = t.rows[0].cells[0]; c.text=""
    p = c.paragraphs[0]; p.paragraph_format.space_before=Pt(4); p.paragraph_format.space_after=Pt(4)
    add_runs(p, text, size=10, color=INK)
    shade(c, PANEL); borders(c, left=(24,color))
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

# footer page numbers
def add_page_numbers():
    for s in doc.sections:
        p = s.footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(); r.font.size=Pt(8); r.font.color.rgb=RGBColor.from_string(INK3)
        fld = OxmlElement("w:fldSimple"); fld.set(qn("w:instr"), "PAGE")
        p._p.append(fld)

exec(open(os.path.join(ROOT,"tools","docx_content.py")).read())
add_page_numbers()
doc.save(OUT)
print(f"wrote {OUT}  ({os.path.getsize(OUT)/1024:.0f} KB)")
